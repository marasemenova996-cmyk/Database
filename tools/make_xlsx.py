"""Собирает scores.csv в Excel-файл database.xlsx: один лист «Таблица», шкала и критерии оценки — справа от таблицы."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

csv.field_size_limit(10**9)
F = Font(name="Arial", size=10)
H = Font(name="Arial", size=10, bold=True)
LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
HF = PatternFill("solid", fgColor="DDE4EE")
LF = PatternFill("solid", fgColor="F3F5F8")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

TOPIC_RULES = [
    ("Россия", "−: Россия как противник, агрессор, санкции против неё; +: переговоры, жесты сближения"),
    ("Украинский кризис", "+: поддержка Украины, помощь, мирное урегулирование; −: давление на Украину"),
    ("Китай", "−2: угроза/противник; −1: конкурент; +1: партнёр по сделке, переговоры"),
    ("Изменение климата", "всегда 0, позиция описывается словами в примечании (напр. «выход из Парижского соглашения»)"),
    ("Миграция", "−3: «вторжение», чрезвычайное положение; −2: нелегальная миграция как угроза; −1: мягкое упоминание; +1: полезная легальная миграция"),
    ("Глобализация", "−: критика «глобализма», международных обязательств; +: поддержка"),
    ("НАТО", "+2: поддержка, похвала; +1: требование делить расходы; −1: «они нам не нужны»"),
    ("ЕС", "+: партнёрство; −: критика"),
    ("ООН", "−2: выход из организаций, резкая критика; −1: критика; +1: работа в органах ООН"),
    ("Международная торговля", "−1: пошлины, «нечестная торговля», дефицит; +1: торговые сделки, открытие рынков"),
    ("Демократии в мире", "+: продвижение демократии и свобод за рубежом; −: отказ от этого"),
    ("Военные интервенции", "+2: одобрение операций и ударов США; −: отказ, сворачивание"),
    ("Либеральная повестка", "−2: против абортов, «гендерной идеологии», DEI; +: поддержка (с примечанием)"),
    ("Израиль", "+2: поддержка; +3: посольство в Иерусалиме; +1: участие в мирной сделке"),
    ("Многополярный мир", "+: признание баланса сил/многополярности; −: отрицание"),
]

def write_table(ws, rows, start, widths, numeric=(), links=True):
    head = list(rows[0].keys())
    for j, k in enumerate(head, 1):
        c = ws.cell(start, j, k)
        c.font, c.fill, c.alignment = H, HF, CENTER
    for i, r in enumerate(rows, start + 1):
        for j, k in enumerate(head, 1):
            v = r[k]
            if j - 1 in numeric and v.lstrip("-").isdigit(): v = int(v)
            c = ws.cell(i, j, v)
            c.font = F
            c.alignment = CENTER if j - 1 in numeric else WRAP
            if links and j == 1 and v.startswith("http"):
                c.hyperlink = v; c.font = LINK
    for j, w in enumerate(widths(head), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[start].height = 75
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(head))}{start + len(rows)}"
    return head

scores = list(csv.DictReader(open("scores.csv", encoding="utf-8")))

wb = Workbook()
ws = wb.active
ws.title = "Таблица"

# Таблица с первой строки, легенда (шкала и критерии) — сбоку справа от таблицы
start = 1
head = write_table(ws, scores, start, lambda h: [45, 12, 16, 22, 16, 11, 26, 12] + [12] * 15 + [40],
                   numeric=set(range(7, 23)) | {2})
ws.freeze_panes = "B2"
legend = [("Как читать таблицу", None),
          ("Шкала оценок", "−3 резко негативно\n−2 негативно\n−1 скорее негативно\n0 нейтрально / упомянуто без оценки\n"
                           "+1 скорее позитивно\n+2 позитивно\n+3 резко позитивно\nпустая клетка — тема не затронута"),
          ("Директивность", "1 — указ, прокламация, меморандум, директива, соглашение, стратегия; 0 — заявления, речи, релизы."),
          ("Отбор", "в таблицу входят только документы, где оценена хотя бы одна тема."),
          ("Критерии по темам", None)] + TOPIC_RULES
lc = len(head) + 2                       # одна пустая колонка между таблицей и легендой
ws.column_dimensions[get_column_letter(lc)].width = 24
ws.column_dimensions[get_column_letter(lc + 1)].width = 60
for r, (name, text) in enumerate(legend, 1):
    a = ws.cell(r, lc, name)
    if text is None:
        a.font = Font(name="Arial", size=12 if r == 1 else 10, bold=True)
        if r > 1: a.fill = HF; ws.cell(r, lc + 1).fill = HF
        a.alignment = WRAP
        continue
    a.font, a.fill, a.alignment = H, LF, WRAP
    b = ws.cell(r, lc + 1, text); b.font, b.fill, b.alignment = F, LF, WRAP
    lines = sum(1 + len(part) // 70 for part in text.split("\n"))
    ws.row_dimensions[r].height = max(30, 13 * lines)
wb.save("database.xlsx")
print("database.xlsx:", len(scores), "строк")
